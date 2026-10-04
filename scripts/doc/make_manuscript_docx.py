#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build  submission docx files from manuscript_.md and supplementary/supplementary.md.
Figures (R/ggplot2 300-dpi PNG) are embedded at their legend positions."""
import os, re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE = "/Users/zuoxianbo/Desktop/SCI论文/Sandy_Nature/outputs/analysis_base"
FIG = os.path.join(BASE, "figures_")

MAINMAP = {
    "Figure 1": "fig1_transport_system.png",
    "Figure 2": "fig2_external_transport.png",
    "Figure 3": "fig3_calibration_pathways.png",
    "Figure 4": "fig4_ablation_constraint.png",
    "Figure 5": "fig5_benchmark_strategies.png",
    "Figure 6": "fig6_joint_boundary.png",
}
EDMAP = {f"Extended Data Fig. {i}": f"ed{i}_"
         for i in range(1, 11)}
EDFILES = {
    "Extended Data Fig. 1": "ed1_positive_control.png",
    "Extended Data Fig. 2": "ed2_cohort_ledger.png",
    "Extended Data Fig. 3": "ed3_loco_sensitivity.png",
    "Extended Data Fig. 4": "ed4_ablation_matrix.png",
    "Extended Data Fig. 5": "ed5_identity_recovery.png",
    "Extended Data Fig. 6": "ed6_threshold_sensitivity.png",
    "Extended Data Fig. 7": "ed7_clinical_casemix.png",
    "Extended Data Fig. 8": "ed8_toxicity_stress.png",
    "Extended Data Fig. 9": "ed9_precision_simulation.png",
    "Extended Data Fig. 10": "ed10_analysis_freeze.png",
}

def add_runs(p, s):
    for part in re.split(r'(\*\*[^*]+\*\*)', s):
        if part == "":
            continue
        if len(part) >= 4 and part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            p.add_run(part)

def add_md_table(doc, rows):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=1, cols=ncol)
    try:
        t.style = "Table Grid"
    except Exception:
        pass
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, c in enumerate(rows[0]):
        hdr[i].text = c
        for para in hdr[i].paragraphs:
            for r in para.runs:
                r.bold = True; r.font.size = Pt(8)
    for row in rows[1:]:
        cells = t.add_row().cells
        for i, c in enumerate(row):
            cells[i].text = c
            for para in cells[i].paragraphs:
                for r in para.runs:
                    r.font.size = Pt(8)

def convert(md_path, out_path, figs=True):
    text = open(md_path, encoding="utf-8").read()
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(11)
    lines = text.split("\n")
    n_img = 0
    i, N = 0, len(lines)
    while i < N:
        ln = lines[i]; s = ln.strip()
        if s.startswith("|") and i + 1 < N and re.match(r'^\s*\|[\s:|-]+\|\s*$', lines[i+1].strip()):
            tbl = []
            while i < N and lines[i].strip().startswith("|"):
                tbl.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            if len(tbl) >= 3:
                add_md_table(doc, [tbl[0]] + tbl[2:])
                doc.add_paragraph()
            continue
        if ln.startswith("# "):
            p = doc.add_paragraph()
            r = p.add_run(ln[2:].strip()); r.bold = True; r.font.size = Pt(14)
            i += 1; continue
        if ln.startswith("## "):
            doc.add_heading(ln[3:].strip(), level=1)
            i += 1; continue
        if ln.startswith("### "):
            doc.add_heading(ln[4:].strip(), level=2)
            i += 1; continue
        # main figure legends -> embed above the legend text
        mmain = re.match(r'^Figure (\d+) \|', s)
        med = re.match(r'^Extended Data Fig\. (\d+) \|', s)
        if figs and (mmain or med):
            tok = (f"Figure {mmain.group(1)}" if mmain
                   else f"Extended Data Fig. {med.group(1)}")
            fn = MAINMAP.get(tok) or EDFILES.get(tok)
            if fn:
                fp = os.path.join(FIG, fn)
                if os.path.exists(fp):
                    doc.add_picture(fp, width=Inches(6.0))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    n_img += 1
            p = doc.add_paragraph()
            r = p.add_run(s); r.font.size = Pt(9)
            i += 1; continue
        if s == "" or s == "---":
            i += 1; continue
        p = doc.add_paragraph()
        add_runs(p, s)
        i += 1
    doc.save(out_path)
    print("DOCX WRITTEN:", out_path, os.path.getsize(out_path), "bytes;", n_img, "images")

convert(os.path.join(BASE, "manuscript", "manuscript_.md"),
        os.path.join(BASE, "_manuscript_NatureCancer.docx"), figs=True)
convert(os.path.join(BASE, "supplementary/supplementary.md"),
        os.path.join(BASE, "_supplementary_NatureCancer.docx"), figs=False)
